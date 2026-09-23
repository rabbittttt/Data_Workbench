from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import json
from typing import Any, Iterable
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
PARSER_VERSION = 'hierarchy-time-v2'
OVERRIDES = Path(__file__).resolve().parent / 'output_file' / 'recognition.json'


@dataclass
class ParsedTable:
    file_path: Path
    sheet_name: str
    table_name: str
    range_ref: str
    dataframe: pd.DataFrame
    header_row: int


def scan_excel_files(folder: str | Path) -> list[Path]:
    root = Path(folder).expanduser().resolve()
    if not root.exists() or not root.is_dir():
        raise ValueError(f"数据目录不存在：{root}")
    return sorted(
        path
        for path in root.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_SUFFIXES
        and not path.name.startswith("~$")
    )


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = re.sub(r"\s+", " ", value).strip()
        return value or None
    return value


def _row_nonempty(row: list[Any]) -> int:
    return sum(not _blank(value) for value in row)


def _active_row_bands(matrix: list[list[Any]]) -> list[tuple[int, int]]:
    """Split a sheet on blank rows while retaining one-row titles."""
    active = [_row_nonempty(row) > 0 for row in matrix]
    bands: list[tuple[int, int]] = []
    start: int | None = None
    blank_run = 0
    for index, is_active in enumerate(active):
        if is_active:
            if start is None:
                start = index
            blank_run = 0
        elif start is not None:
            blank_run += 1
            if blank_run >= 1:
                bands.append((start, index - blank_run))
                start = None
                blank_run = 0
    if start is not None:
        bands.append((start, len(matrix) - 1))
    return bands


def _column_bands(matrix: list[list[Any]], top: int, bottom: int) -> list[tuple[int, int]]:
    max_cols = max((len(row) for row in matrix[top : bottom + 1]), default=0)
    active = []
    for col in range(max_cols):
        count = sum(
            col < len(matrix[row]) and not _blank(matrix[row][col])
            for row in range(top, bottom + 1)
        )
        active.append(count > 0)
    bands: list[tuple[int, int]] = []
    start: int | None = None
    for index, is_active in enumerate(active + [False]):
        if is_active and start is None:
            start = index
        elif not is_active and start is not None:
            bands.append((start, index - 1))
            start = None
    return bands


def _header_score(rows: list[list[Any]], candidate: int) -> float:
    row = rows[candidate]
    values = [value for value in row if not _blank(value)]
    if len(values) < 2:
        return -1
    strings = sum(isinstance(value, str) for value in values)
    unique = len({str(value).strip().lower() for value in values}) / len(values)
    below = rows[candidate + 1 : candidate + 5]
    data_rows = sum(_row_nonempty(item) >= 2 for item in below)
    if data_rows == 0:
        return -1
    string_ratio = strings / len(values)
    next_values = [value for item in below for value in item if not _blank(value)]
    next_string_ratio = (
        sum(isinstance(value, str) for value in next_values) / len(next_values)
        if next_values
        else 1
    )
    type_contrast = max(0.0, string_ratio - next_string_ratio)
    periods = sum(isinstance(v, (date, datetime)) or bool(re.fullmatch(
        r'(?:\d{4}[-/.年])?\d{1,2}[-/.月]\d{1,2}日?|\d{1,2}月|\d{2,4}WK\d{1,2}|D\d+', str(v), re.I)) for v in values)
    return string_ratio * 3 + unique + type_contrast * 2 + min(data_rows, 3) * 0.4 + (5 if periods >= 2 else 0)


def _fill_merged_labels(worksheet, matrix, header_row, left, right, bottom):
    """Fill only actual vertical merges below this header; ordinary blanks stay blank."""
    for merged in worksheet.merged_cells.ranges:
        if merged.min_col != merged.max_col or not (left <= merged.min_col <= right):
            continue
        if merged.min_row <= header_row or merged.min_row > bottom:
            continue
        value = matrix[merged.min_row - 1][merged.min_col - 1]
        for row in range(merged.min_row, min(merged.max_row, bottom) + 1):
            matrix[row - 1][merged.min_col - 1] = value


def _make_headers(values: Iterable[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if not _blank(value) else f"字段{index}"
        base = re.sub(r"\s+", " ", base)
        candidate = base
        count = 1
        while candidate.casefold() in seen or candidate.startswith('_'):
            count += 1
            candidate = f"{base.lstrip('_')}_{count}"
        seen[candidate.casefold()] = 1
        headers.append(candidate)
    return headers


def _table_title(rows: list[list[Any]], header_index: int, fallback: str) -> str:
    for index in range(header_index - 1, -1, -1):
        values = [str(v).strip() for v in rows[index] if not _blank(v)]
        if 1 <= len(values) <= 3:
            title = " ".join(values)
            if len(title) <= 80:
                return title
    return fallback


def discover_tables(path: str | Path, rules=None) -> list[ParsedTable]:
    file_path = Path(path).resolve()
    workbook = load_workbook(file_path, data_only=True, read_only=False)
    discovered: list[ParsedTable] = []
    try:
        for worksheet in workbook.worksheets:
            matrix = [[_clean(value) for value in row] for row in worksheet.iter_rows(values_only=True)]
            if not matrix:
                continue
            for band_number, (top, bottom) in enumerate(_active_row_bands(matrix), start=1):
                if bottom - top < 1:
                    continue
                for left, right in _column_bands(matrix, top, bottom):
                    if right - left < 1:
                        continue
                    rows = [row[left : right + 1] for row in matrix[top : bottom + 1]]
                    search_limit = min(len(rows) - 1, 12)
                    scores = [(_header_score(rows, idx), idx) for idx in range(search_limit)]
                    score, header_idx = max(scores, key=lambda pair: (pair[0], -pair[1]), default=(-1, -1))
                    if score < 0:
                        continue
                    _fill_merged_labels(worksheet, matrix, top + header_idx + 1, left + 1, right + 1, bottom + 1)
                    rows = [row[left : right + 1] for row in matrix[top : bottom + 1]]
                    headers = _make_headers(rows[header_idx])
                    body = rows[header_idx + 1 :]
                    body = [row for row in body if _row_nonempty(row) >= 1]
                    if not body:
                        continue
                    frame = pd.DataFrame(body, columns=headers)
                    frame = frame.dropna(axis=1, how="all").dropna(axis=0, how="all")
                    if frame.empty or frame.shape[1] < 2:
                        continue
                    frame = frame.reset_index(drop=True)
                    fallback = f"{worksheet.title} 表{band_number}"
                    title = _table_title(rows, header_idx, fallback)
                    if title == fallback and top > 0:
                        # A title may be separated from its header by blank rows.
                        for prior in range(top - 1, max(-1, top - 5), -1):
                            previous = [v for v in matrix[prior][left:right+1] if not _blank(v)]
                            if not previous:
                                continue
                            if len(previous) == 1 and isinstance(previous[0], str) and len(previous[0]) <= 80:
                                title = previous[0]
                            break
                    start_row = top + header_idx + 1
                    end_row = top + header_idx + 1 + len(body)
                    range_ref = f"{get_column_letter(left + 1)}{start_row}:{get_column_letter(left + len(headers))}{end_row}"
                    discovered.append(
                        ParsedTable(file_path, worksheet.title, title, range_ref, frame, start_row)
                    )
    finally:
        workbook.close()
    rules = rules if rules is not None else (json.loads(OVERRIDES.read_text(encoding='utf-8')) if OVERRIDES.exists() else {})
    relevant = [rule for rule in rules.values() if Path(rule['file']).resolve() == file_path]
    if relevant:
        from openpyxl.utils.cell import range_boundaries
        book = load_workbook(file_path, read_only=False, data_only=True)
        try:
            for rule in relevant:
                left, top, right, bottom = range_boundaries(rule['range'])
                sheet = book[rule['sheet']]
                if bottom > sheet.max_row or right > sheet.max_column:
                    raise ValueError('指定区域超过工作表有效范围')
                matrix = [[_clean(v) for v in row] for row in sheet.iter_rows(values_only=True)]
                _fill_merged_labels(sheet, matrix, top, left, right, bottom)
                values = [row[left - 1:right] for row in matrix[top - 1:bottom]]
                frame = pd.DataFrame(values[1:], columns=_make_headers(values[0]))
                if frame.empty:
                    raise ValueError('指定区域没有数据')
                original = rule.get('original', rule['range']).split(':')[0]
                discovered = [t for t in discovered if not (t.sheet_name == rule['sheet'] and t.range_ref.split(':')[0] in (original, rule['range'].split(':')[0]))]
                discovered.append(ParsedTable(file_path, rule['sheet'], rule['name'], rule['range'], frame, top))
        finally:
            book.close()
    return discovered


def infer_series_type(series: pd.Series) -> str:
    values = series.dropna()
    if values.empty:
        return "empty"
    if pd.api.types.is_datetime64_any_dtype(values):
        return "date"
    numeric = pd.to_numeric(values, errors="coerce")
    if numeric.notna().mean() >= 0.8:
        return "number"
    header = str(series.name).lower()
    date_hint = any(word in header for word in ("日期", "时间", "年月", "date", "time", "month"))
    has_date_values = values.map(lambda value: isinstance(value, (date, datetime, pd.Timestamp))).any()
    if date_hint or has_date_values:
        # Date columns may mix Excel dates, ISO strings, and localized text dates.
        # ``format='mixed'`` keeps per-value inference explicit on pandas 2.2+
        # and avoids the warning about falling back to dateutil.
        parsed_dates = pd.to_datetime(values, format="mixed", errors="coerce")
        if parsed_dates.notna().mean() >= 0.8:
            return "date"
    return "text"
