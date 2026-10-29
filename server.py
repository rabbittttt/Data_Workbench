"""Local data workspace. Run: python server.py (http://127.0.0.1:8765)."""
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from http.client import HTTPConnection, HTTPException
from pathlib import Path
from urllib.parse import urlparse, parse_qs, parse_qsl, urlencode, unquote
from dataclasses import asdict
from html import escape, unescape
from io import BytesIO
import json
import os
import hashlib
import re
import argparse
import socket
import threading
import time
import tempfile
import webbrowser
import pandas as pd
from openpyxl import load_workbook
from warehouse import Warehouse

ROOT = Path(__file__).resolve().parent
STORE = Warehouse(ROOT / 'output_file' / 'warehouse.db')
LOCK = threading.RLock()
SYNC_LOCK = threading.Lock()
STATUS = {'busy': False, 'errors': [], 'last_sync': None}
CONFIG = ROOT / 'output_file' / 'workspace.json'
MAX_IMPORT_BYTES = 256 * 1024 * 1024


def config():
    return json.loads(CONFIG.read_text(encoding='utf-8')) if CONFIG.exists() else {'folder': str(ROOT / 'input_file'), 'auto': True}


def claim_sync():
    with SYNC_LOCK:
        if STATUS['busy']:
            STATUS['pending'] = True
            return False
        STATUS['busy'] = True
        return True


def sync_worker():
    while True:
        try:
            with LOCK:
                result = STORE.sync(config()['folder'])
            STATUS.update(asdict(result), last_sync=time.strftime('%Y-%m-%d %H:%M:%S'))
        except Exception as exc:
            STATUS['errors'] = [str(exc)]
        with SYNC_LOCK:
            # A request made during a scan must include files added after it began.
            if STATUS.pop('pending', False):
                continue
            STATUS['busy'] = False
            return


def sync():
    if claim_sync():
        sync_worker()


def request_sync():
    if claim_sync():
        threading.Thread(target=sync_worker, daemon=True).start()


def import_workbook(stream, name, size, expected_folder):
    """Store the original bytes; incomplete uploads never become scanned workbooks."""
    if (not name or len(name) > 240 or name.endswith((' ', '.'))
            or re.search(r'[<>:"/\\|?*\x00-\x1f]', name)
            or name.startswith('~$')
            or re.fullmatch(r'CON|PRN|AUX|NUL|COM[1-9¹²³]|LPT[1-9¹²³]', name.split('.')[0], re.I)):
        raise ValueError('文件名无效，请去掉路径、特殊字符或 Excel 临时文件前缀')
    if Path(name).suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise ValueError('仅支持 .xlsx 和 .xlsm；旧版 .xls 请先另存为 .xlsx')
    if size <= 0:
        raise ValueError('文件为空或缺少文件长度')
    folder = Path(config()['folder']).resolve()
    if not expected_folder or Path(expected_folder).resolve() != folder:
        raise ValueError('数据源目录已变更，请重新打开导入窗口')
    if folder == (ROOT / 'input_file').resolve():
        folder.mkdir(exist_ok=True)
    if not folder.is_dir():
        raise ValueError('数据源目录不存在，请在“管理数据源”中重新选择')
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=folder, prefix='.import-', suffix='.part', delete=False) as output:
            temporary = Path(output.name)
            remaining = size
            while remaining:
                chunk = stream.read(min(1024 * 1024, remaining))
                if not chunk:
                    raise ValueError('上传中断，文件未保存，请重新选择后导入')
                output.write(chunk)
                remaining -= len(chunk)
        try:
            with temporary.open('rb') as source:
                book = load_workbook(source, read_only=True, data_only=True)
                book.close()
        except Exception as exc:
            raise ValueError('无法读取此 Excel，文件可能损坏或已加密，请用 Excel 打开后另存再导入') from exc
        with LOCK:
            if Path(config()['folder']).resolve() != folder:
                raise ValueError('上传期间数据源目录已变更，请重新打开导入窗口')
            destination = folder / name
            number = 1
            while True:
                try:
                    # Windows rename and POSIX hard-link creation both refuse an
                    # existing destination, including concurrent same-name imports.
                    if os.name == 'nt':
                        temporary.rename(destination)
                    else:
                        os.link(temporary, destination)
                    break
                except FileExistsError:
                    destination = folder / f'{Path(name).stem} ({number}){Path(name).suffix}'
                    number += 1
        return {'saved': True, 'name': destination.name, 'renamed': destination.name != name}
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def catalog():
    records = STORE.catalog().to_dict('records')
    folder = Path(config()['folder']).resolve()
    for row in records:
        row['fields'] = json.loads(row.pop('columns_json'))
        path = Path(row['file_path'])
        row['relative_path'] = str(path.relative_to(folder)) if path.is_relative_to(folder) else str(path)
        row['archived'] = not path.exists() or not path.is_relative_to(folder)
        row['issues'] = []
        if any(f['name'].startswith('字段') for f in row['fields']):
            row['issues'].append('存在未命名字段，建议核对表头')
        if any(f['non_null'] < row['row_count']*.5 for f in row['fields']):
            row['issues'].append('部分字段缺失超过一半')
        if row['row_count'] <= 4:
            row['issues'].append('小型数据区域，可能是图表辅助区')
    return records


def data(dataset_id):
    frame = STORE.load(dataset_id)
    # JSON serializer preserves nulls and ISO dates; no NaN reaches the browser.
    return {'columns': list(frame.columns), 'rows': json.loads(frame.to_json(orient='records', date_format='iso', force_ascii=False))}


def workbook_files():
    folder = Path(config()['folder']).resolve()
    return {hashlib.sha256(str(p.resolve()).encode()).hexdigest()[:24]: p.resolve()
            for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in {'.xlsx', '.xlsm'}
            and not p.name.startswith('~$') and p.resolve().is_relative_to(folder)}


def workbook_list():
    folder = Path(config()['folder']).resolve()
    files = []
    for key, path in sorted(workbook_files().items(), key=lambda item: str(item[1])):
        try:
            stat = path.stat()
        except FileNotFoundError:
            continue
        files.append(dict(id=key, name=str(path.relative_to(folder)), path=str(path),
                          signature=f'{stat.st_mtime_ns}:{stat.st_size}'))
    return files


def workbook_info(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return {'sheets': [sheet.title for sheet in book.worksheets]}
    finally:
        book.close()


def original_sheet(path, sheet_name=None, start=1, column=1):
    """Read source cells independently of inferred tables. Never save the workbook."""
    book = load_workbook(path, data_only=True)
    try:
        sheets = [sheet.title for sheet in book.worksheets]
        missing_sheet = sheet_name if sheet_name and sheet_name not in sheets else None
        if not sheets:
            raise ValueError('此文件没有可读取的单元格工作表')
        sheet = book[sheet_name] if sheet_name and not missing_sheet else book.worksheets[0]
        if missing_sheet:
            start = column = 1
        start = min(max(1, start), sheet.max_row)
        column = min(max(1, column), sheet.max_column)
        bottom, right = min(start+79, sheet.max_row), min(column+29, sheet.max_column)
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(
            min_row=start, max_row=bottom, min_col=column, max_col=right)]
        merges = [list(r.bounds) for r in sheet.merged_cells.ranges
                  if r.min_row <= bottom and r.max_row >= start and r.min_col <= right and r.max_col >= column]
        return dict(sheets=sheets, sheet=sheet.title, rows=rows, merges=merges,
                    start=start, column=column, max_row=sheet.max_row, max_col=sheet.max_column,
                    missing_sheet=missing_sheet,
                    warning=f'原工作表“{missing_sheet}”已不存在，已显示“{sheet.title}”。请重新选择当前工作表；旧图表未改动。' if missing_sheet else '')
    finally:
        book.close()


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT / 'web'), **kwargs)

    def versioned_html(self, path, body):
        root = (ROOT / 'web').resolve()
        versions = {}

        def version_asset(match):
            url = urlparse(unescape(match['url']))
            if url.scheme or url.netloc or not url.path:
                return match[0]
            base = root if url.path.startswith('/') else path.parent
            asset = (base / unquote(url.path).lstrip('/')).resolve()
            if not asset.is_relative_to(root) or not asset.is_file():
                return match[0]
            if asset not in versions:
                versions[asset] = hashlib.sha256(asset.read_bytes()).hexdigest()[:16]
            query = [(key, value) for key, value in parse_qsl(url.query) if key != 'v']
            query.append(('v', versions[asset]))
            versioned = url._replace(query=urlencode(query)).geturl()
            return match['prefix'] + escape(versioned, quote=True) + match['quote']

        return re.sub(r'(?P<prefix>\b(?:src|href)=(?P<quote>[\"\x27]))(?P<url>[^\"\x27]+)(?P=quote)',
                      version_asset, body.decode('utf-8')).encode('utf-8')

    def send_head(self):
        path = Path(self.translate_path(self.path))
        if path.is_dir():
            if not urlparse(self.path).path.endswith('/'):
                return super().send_head()
            path = path / 'index.html'
        if not path.is_file():
            return super().send_head()
        try:
            body = path.read_bytes()
            if path.suffix.lower() == '.html':
                body = self.versioned_html(path, body)
        except OSError:
            self.send_error(404, 'File not found')
            return None
        # Content validators stay correct after Git pulls even when timestamps match.
        # Deliberately ignore If-Modified-Since; it cannot validate changed bytes.
        etag = '"' + hashlib.sha256(body).hexdigest() + '"'
        candidates = [value.strip().removeprefix('W/')
                      for value in self.headers.get('If-None-Match', '').split(',')]
        unchanged = etag in candidates or '*' in candidates
        self.send_response(304 if unchanged else 200)
        self.send_header('Cache-Control', 'no-cache')
        self.send_header('ETag', etag)
        if not unchanged:
            content_type = self.guess_type(str(path))
            if content_type == 'text/html':
                content_type += '; charset=utf-8'
            self.send_header('Content-Type', content_type)
            self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        return None if unchanged else BytesIO(body)

    def send_json(self, payload, code=200):
        body = json.dumps(payload, ensure_ascii=False, default=str).encode('utf-8')
        try:
            self.send_response(code)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Cache-Control', 'no-store')
            self.send_header('Content-Length', str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
            # The client is gone; do not send another response on this socket.
            self.close_connection = True

    def do_GET(self):
        parts = urlparse(self.path)
        query = parse_qs(parts.query)
        try:
            if parts.path == '/api/catalog':
                return self.send_json({'tables': catalog(), 'files': workbook_list(), 'status': STATUS, 'config': config()})
            if parts.path == '/api/table':
                return self.send_json(data(query['id'][0]))
            if parts.path == '/api/workbooks':
                return self.send_json({'files': workbook_list()})
            if parts.path == '/api/workbook':
                path = workbook_files().get(query['id'][0])
                if path is None:
                    raise ValueError('文件已移走，请重新选择数据文件')
                return self.send_json(workbook_info(path))
            if parts.path == '/api/sheet':
                path = workbook_files().get(query['id'][0])
                if path is None:
                    raise ValueError('文件已移走，请重新选择数据文件')
                return self.send_json(original_sheet(path, query.get('sheet', [None])[0],
                                      int(query.get('start', ['1'])[0]), int(query.get('column', ['1'])[0])))
            if parts.path == '/api/search':
                words = query.get('q', [''])[0].casefold().split()
                hits = []
                if words:
                    for table in catalog():
                        context = ' '.join([table['file_name'], table['sheet_name'], table['table_name'], *[f['name'] for f in table['fields']]]).casefold()
                        frame = STORE.load(table['id'])
                        matched = []
                        for row in frame.to_dict('records'):
                            if all(w in context + ' ' + ' '.join(str(v) for v in row.values()).casefold() for w in words):
                                matched.append(row)
                                if len(matched) >= 3:
                                    break
                        if matched or all(w in context for w in words):
                            hits.append({'table': table, 'rows': json.loads(pd.DataFrame(matched).to_json(orient='records', date_format='iso'))})
                return self.send_json({'hits': hits})
            if parts.path == '/api/raw':
                table = next(t for t in catalog() if t['id'] == query['id'][0])
                book = load_workbook(table['file_path'], read_only=True, data_only=True)
                try:
                    sheet = book[table['sheet_name']]
                    start = max(1, int(query.get('start', ['1'])[0]))
                    rows = list(sheet.iter_rows(min_row=start, max_row=min(start+79, sheet.max_row), max_col=min(sheet.max_column, 60), values_only=True))
                    return self.send_json({'rows': rows, 'start': start, 'max_row': sheet.max_row, 'max_col': sheet.max_column})
                finally:
                    book.close()
            if parts.path.startswith('/api/'):
                return self.send_json({'error': '接口不存在'}, 404)
            return super().do_GET()
        except Exception as exc:
            self.send_json({'error': str(exc)}, 400)

    def do_POST(self):
        # Non-simple content types + same origin; listen exclusively on loopback.
        if self.headers.get('Origin') and self.headers['Origin'] != 'http://' + self.headers.get('Host', ''):
            return self.send_json({'error': '来源不匹配'}, 403)
        if urlparse(self.path).path == '/api/import':
            try:
                if self.headers.get('Content-Type', '').split(';')[0] != 'application/octet-stream':
                    raise ValueError('需要文件二进制请求')
                if self.headers.get('Transfer-Encoding'):
                    raise ValueError('需要明确的文件长度')
                size = int(self.headers.get('Content-Length', 0))
                if size > MAX_IMPORT_BYTES:
                    return self.send_json({'error': '单个文件不能超过 256 MB'}, 413)
                query = parse_qs(urlparse(self.path).query)
                self.connection.settimeout(120)
                return self.send_json(import_workbook(self.rfile, query.get('filename', [''])[0], size,
                                                      query.get('folder', [''])[0]), 201)
            except (TimeoutError, ConnectionError):
                return self.send_json({'error': '上传中断或超时，请检查连接后重新导入'}, 400)
            except Exception as exc:
                return self.send_json({'error': str(exc)}, 400)
        if 'application/json' not in self.headers.get('Content-Type', ''):
            return self.send_json({'error': '需要 JSON 请求'}, 400)
        try:
            payload = json.loads(self.rfile.read(min(int(self.headers.get('Content-Length', 0)), 65536)) or b'{}')
            if self.path == '/api/sync':
                request_sync()
                return self.send_json({'started': True})
            if self.path == '/api/config':
                folder = Path(payload['folder']).expanduser().resolve()
                if not folder.is_dir():
                    raise ValueError('文件夹不存在')
                with LOCK:
                    CONFIG.write_text(json.dumps({'folder': str(folder), 'auto': bool(payload.get('auto', True))}, ensure_ascii=False), encoding='utf-8')
                request_sync()
                return self.send_json({'saved': True})
            if self.path == '/api/recognition':
                from excel_parser import discover_tables, OVERRIDES
                table = next(t for t in catalog() if t['id'] == payload['id'])
                from openpyxl.utils.cell import range_boundaries
                left, top, right, bottom = range_boundaries(payload['range'])
                if not all([left, top, right, bottom]) or bottom <= top or right < left:
                    raise ValueError('区域需包含一行表头及至少一行数据，例如 A2:F80')
                with LOCK:
                    rules = json.loads(OVERRIDES.read_text(encoding='utf-8')) if OVERRIDES.exists() else {}
                    key = table['file_path'] + '|' + table['sheet_name'] + '|' + table['range_ref'].split(':')[0]
                    rules[key] = {'file': table['file_path'], 'sheet': table['sheet_name'], 'range': payload['range'], 'name': payload.get('name') or table['table_name'], 'original': table['range_ref']}
                    # Validate the entire workbook before persisting the rule.
                    parsed = discover_tables(table['file_path'], rules=rules)
                    STORE._replace_file(Path(table['file_path']), parsed, '')
                    OVERRIDES.write_text(json.dumps(rules, ensure_ascii=False, indent=2), encoding='utf-8')
                return self.send_json({'saved': True})
            self.send_json({'error': '接口不存在'}, 404)
        except Exception as exc:
            self.send_json({'error': str(exc)}, 400)


def watch():
    while True:
        if config().get('auto', True):
            sync()
        time.sleep(60)


WORKBENCH_URL = 'http://127.0.0.1:8765'


class WorkbenchHTTPServer(ThreadingHTTPServer):
    def server_bind(self):
        if hasattr(socket, 'SO_EXCLUSIVEADDRUSE'):
            # Windows must not let a second launcher share the same listening port.
            self.allow_reuse_address = False
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()


def open_workbench():
    try:
        if webbrowser.open(WORKBENCH_URL, new=2):
            return
    except webbrowser.Error:
        pass
    print(f'未能自动打开浏览器，请手动访问 {WORKBENCH_URL}', flush=True)


def workbench_running():
    connection = HTTPConnection('127.0.0.1', 8765, timeout=2)
    try:
        connection.request('GET', '/')
        response = connection.getresponse()
        title = '<title>数见 · 经营数据工作台</title>'.encode('utf-8')
        return response.status == 200 and title in response.read(16384)
    except (OSError, HTTPException):
        return False
    finally:
        connection.close()


def run_server(open_browser=False):
    try:
        httpd = WorkbenchHTTPServer(('127.0.0.1', 8765), Handler)
    except OSError:
        if open_browser and workbench_running():
            print('工作台已经运行，正在打开现有页面。', flush=True)
            open_workbench()
            return
        raise
    with httpd:
        threading.Thread(target=watch, daemon=True).start()
        print(f'数据速查工作台: {WORKBENCH_URL}', flush=True)
        if open_browser:
            # The listening socket is ready before the browser can request the page.
            threading.Thread(target=open_workbench, daemon=True).start()
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            pass


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='数见 · 经营数据工作台')
    parser.add_argument('--open-browser', action='store_true', help='服务启动后自动打开浏览器')
    run_server(open_browser=parser.parse_args().open_browser)
