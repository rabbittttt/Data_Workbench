import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from server import original_sheet


class OriginalSheetTests(unittest.TestCase):
    def test_unrecognized_sheets_positions_and_paging(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = '原表'
            sheet['B3'] = '标题'
            sheet.merge_cells('B3:D3')
            sheet['BK101'] = 42
            book.create_sheet('只有说明')['A1'] = '无需识别也能看见'
            book.save(path)
            before = path.read_bytes()
            first = original_sheet(path)
            self.assertEqual(first['rows'][2][1], '标题')
            self.assertEqual(first['merges'], [[2, 3, 4, 3]])
            self.assertIn('只有说明', first['sheets'])
            page = original_sheet(path, '原表', 101, 63)
            self.assertEqual(page['rows'], [[42]])
            self.assertEqual(original_sheet(path, '只有说明')['rows'][0][0], '无需识别也能看见')
            self.assertEqual(before, path.read_bytes())


if __name__ == '__main__':
    unittest.main()
