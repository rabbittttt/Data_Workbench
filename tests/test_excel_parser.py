from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from excel_parser import discover_tables


class ExcelParserTest(unittest.TestCase):
    def test_merged_parent_keeps_children_and_does_not_fill_ordinary_blanks(self):
        from datetime import datetime
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'merged.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['指标', '统计类型', '分类', datetime(2026, 8, 1), datetime(2026, 8, 2)])
            sheet.append(['动力', '占比', '增程', .4, .5])
            sheet.append([None, None, '纯电', .6, .5])
            sheet.append(['座椅', None, '五座', .8, .9])
            sheet.merge_cells('A2:A3')
            sheet.merge_cells('B2:B3')
            book.save(path)
            table = discover_tables(path, rules={})[0]
            self.assertEqual(table.header_row, 1)
            self.assertEqual(table.dataframe['指标'].tolist(), ['动力', '动力', '座椅'])
            self.assertEqual(table.dataframe['分类'].tolist(), ['增程', '纯电', '五座'])
            self.assertTrue(table.dataframe['统计类型'].isna().iloc[2])
            rules = {'custom': {'file': str(path), 'sheet': sheet.title, 'range': 'A1:E4', 'name': '自定义'}}
            custom = discover_tables(path, rules=rules)[0]
            self.assertEqual(custom.dataframe.iloc[1, 0], '动力')

    def test_discovers_two_tables_in_one_sheet(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "sample.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Sheet1"
            rows = [
                ["2026年销量情况"],
                ["车型", "1月", "2月", "3月"],
                ["M5", 100, 200, 300],
                ["M7", 200, 300, 400],
                [],
                ["版本销量占比"],
                ["版本", "占比"],
                ["Pro", 0.4],
                ["Max", 0.5],
            ]
            for row in rows:
                sheet.append(row)
            book.save(path)

            tables = discover_tables(path)

            self.assertEqual(len(tables), 2)
            self.assertEqual(tables[0].table_name, "2026年销量情况")
            self.assertEqual(list(tables[0].dataframe.columns), ["车型", "1月", "2月", "3月"])
            self.assertEqual(tables[1].table_name, "版本销量占比")
            self.assertEqual(tables[1].dataframe.shape, (2, 2))


if __name__ == "__main__":
    unittest.main()
