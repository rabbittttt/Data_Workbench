"""Exercise live HTTP endpoints against a disposable workbook, never business files."""
import json
import tempfile
import threading
import unittest
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen
from unittest.mock import patch
from http.server import ThreadingHTTPServer

from openpyxl import Workbook
import server
from warehouse import Warehouse


class WorkbookRefreshTests(unittest.TestCase):
    def test_source_and_imported_data_follow_saved_workbook_changes(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / 'test.xlsx'
            book = Workbook()
            orders = book.active
            orders.title = '订单'
            orders.append(['车型', '销量'])
            orders.append(['M8', 100])
            orders.append(['M9', 200])
            cache = book.create_sheet('_测试缓存')
            cache.sheet_state = 'veryHidden'
            cache.append(['车型', '销量'])
            cache.append(['M8', 50])
            cache.append(['M9', 60])
            book.save(path)
            store = Warehouse(root / 'test.db')
            status = {'busy': False, 'errors': [], 'last_sync': None}
            with patch.object(server, 'STORE', store), patch.object(server, 'STATUS', status), \
                    patch.object(server, 'config', return_value={'folder': str(root), 'auto': True}):
                httpd = ThreadingHTTPServer(('127.0.0.1', 0), server.Handler)
                worker = threading.Thread(target=httpd.serve_forever, daemon=True)
                worker.start()
                def request(endpoint, **params):
                    url = f'http://127.0.0.1:{httpd.server_port}/api/{endpoint}'
                    if params:
                        url += '?' + urlencode(params)
                    with urlopen(url, timeout=5) as response:
                        return json.load(response)
                try:
                    server.sync()
                    first = request('catalog')
                    file = first['files'][0]
                    self.assertEqual(request('workbook', id=file['id'])['sheets'], ['订单', '_测试缓存'])
                    table = next(t for t in first['tables'] if t['sheet_name'] == '订单')
                    self.assertEqual(request('table', id=table['id'])['rows'][0]['销量'], 100)

                    orders['B2'] = 999
                    orders.append(['M7', 300])
                    orders.title = '订单新版'
                    extra = book.create_sheet('新增数据')
                    extra.append(['车型', '库存'])
                    extra.append(['M8', 7])
                    extra.append(['M9', 8])
                    book.save(path)
                    before = path.read_bytes()
                    self.assertNotEqual(request('workbooks')['files'][0]['signature'], file['signature'])
                    raw = request('sheet', id=file['id'], sheet='订单新版')
                    self.assertEqual(raw['rows'][1][1], 999)
                    self.assertEqual(raw['rows'][-1], ['M7', 300])
                    self.assertEqual(raw['sheets'], ['订单新版', '_测试缓存', '新增数据'])
                    old = request('sheet', id=file['id'], sheet='订单', start=500)
                    self.assertEqual(old['missing_sheet'], '订单')
                    self.assertEqual(old['sheet'], '订单新版')

                    server.sync()
                    updated = request('catalog')
                    self.assertEqual(updated['status']['errors'], [])
                    self.assertEqual({t['sheet_name'] for t in updated['tables']}, {'订单新版', '_测试缓存', '新增数据'})
                    current = next(t for t in updated['tables'] if t['sheet_name'] == '订单新版')
                    rows = request('table', id=current['id'])['rows']
                    self.assertEqual(rows[0]['销量'], 999)
                    self.assertEqual(len(rows), 3)
                    self.assertEqual(path.read_bytes(), before, 'reading and importing never rewrite the workbook')

                    book.remove(extra)
                    book.save(path)
                    server.sync()
                    self.assertEqual(request('workbook', id=file['id'])['sheets'], ['订单新版', '_测试缓存'])
                    self.assertNotIn('新增数据', {t['sheet_name'] for t in request('catalog')['tables']})
                finally:
                    httpd.shutdown()
                    httpd.server_close()
                    worker.join(timeout=5)
                    book.close()


if __name__ == '__main__':
    unittest.main()
